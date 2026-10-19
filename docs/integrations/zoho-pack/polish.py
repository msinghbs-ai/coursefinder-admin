import sys, docx
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=docx.Document(sys.argv[1])
for s in d.sections:
    s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.left_margin=s.right_margin=Cm(2); s.top_margin=s.bottom_margin=Cm(2)
    s.header.paragraphs[0].text="StudySearch Zoho API v2 — Developer Guide"
    s.footer.paragraphs[0].text="zoho-api-v2 · Issue 3 · 19 October 2026 · Pilot"
    for p in (s.header.paragraphs[0],s.footer.paragraphs[0]):
        for r in p.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(0x60,0x60,0x60)
for name in ['Verbatim Char']:
    try: d.styles[name].font.size=Pt(9)
    except KeyError: pass
for name in ['Normal','Body Text','First Paragraph','Compact']:
    try: st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(10)
    except KeyError: pass
for t in d.tables:
    tblPr=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'BFBFBF'); b.append(x)
    old=tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(b)
    hdr=t.rows[0]
    empty=all(not c.text.strip() for c in hdr.cells)
    if hdr.cells[0].text.strip()=='For': continue
    for c in hdr.cells:
        sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'1F3864' if not empty else 'FFFFFF')
        c._tc.get_or_add_tcPr().append(sh)
        for p in c.paragraphs:
            for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    if empty: hdr._tr.getparent().remove(hdr._tr)
d.save(sys.argv[2])
