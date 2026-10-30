#!/usr/bin/env python3
"""Daily New Zealand field report per provider (CF-247, Platform Admin 21 Oct 2026: "Prepare the nz data per uni including all fields
before and at present" and "Produce it daily please").

Input: a file holding the JSON object returned by the query in docs/data-admission/NZ-FIELDS-README.md (it may be wrapped in other text,
for example a saved tool result; the object is found by its "rows" key). It reads pipeline.provider_field_daily (StudySearch Pilot,
read-only): today, the previous day and the start of 21 Oct 2026.
Output: docs/data-admission/nz-fields/NZ-course-data-<date>.xlsx with sheets Read me, NZ summary, Per provider, Long list and, when the
result has "tools" (pipeline.toolset_use_daily, from 22 Oct 2026), Toolset use; and Resolutions & lessons (docs/data-admission/nz-resolutions.json
with the result's "measures", security.nz_resolution_measures_v1).
Usage: python3 scripts/data-admission/build_nz_fields.py result.txt
22 Oct 2026 (Platform Admin 22:56: "prepare Excel sheet like nz to include au and ca as well daily"): the result's "country" (NZ when absent)
picks the country: docs/data-admission/<cc>-fields/<CC>-course-data-<date>.xlsx; the query is in docs/data-admission/COUNTRY-FIELDS-README.md.
31 Oct 2026 (Platform Admin: "include each course with all course field values and scholarship field and entry field"): optional further files,
each one part of the course detail query in COUNTRY-FIELDS-README.md (catalogue, read-only), add the sheets Courses (one row per active course with
every course field value, or why it is not held), Entry requirements (provider entry rules) and Scholarships (active scholarships per provider).
Usage: python3 scripts/data-admission/build_nz_fields.py result.txt [detail-part-0.txt detail-part-1.txt ...]
"""
import html, json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

BASE_DIR = os.path.join(os.path.dirname(__file__), '..', '..', 'docs', 'data-admission')
COUNTRY = {'NZ': 'New Zealand', 'AU': 'Australia', 'CA': 'Canada'}
CC = 'NZ'
CF = [('official_url', 'Official course page'), ('level', 'Study level'), ('field_of_study', 'Field of study'), ('duration', 'Duration'),
      ('delivery', 'Delivery mode'), ('campus', 'Campus linked'), ('campus_city', 'Campus city'), ('international', 'Open to international students'),
      ('provider_tuition', 'Tuition fee'), ('intakes', 'Intakes / start dates'), ('deadline', 'Application deadline'), ('english', 'English requirements'),
      ('academic_entry', 'Academic entry requirements'), ('description', 'Course description')]
PF = [('p_website', 'Website'), ('p_address', 'Street address'), ('p_campuses', 'Campuses'), ('p_type', 'Institution type'),
      ('p_international', 'Enrols international students'), ('p_logo', 'Logo'), ('p_contact', 'Phone or email'), ('p_scholarships', 'Scholarships'),
      ('p_description', 'Description')]
F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True); FH = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HF = PatternFill('solid', fgColor='1F3A5F'); SF = PatternFill('solid', fgColor='DCE6F1')
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
C = Alignment(horizontal='center', vertical='center', wrap_text=True); MID = Alignment(horizontal='center')


def find(obj, depth=0, want=lambda o: 'rows' in o and 'today' in o):
    """Find the report object (by default a dict with "today" and "rows"), possibly JSON-encoded inside strings several times."""
    if depth > 6: return None
    if isinstance(obj, dict):
        if want(obj): return obj
        for v in obj.values():
            f = find(v, depth + 1, want)
            if f: return f
    elif isinstance(obj, list):
        for v in obj:
            f = find(v, depth + 1, want)
            if f: return f
    elif isinstance(obj, str):
        dec = json.JSONDecoder()
        for m in re.finditer(r'[\[{]', obj):
            try:
                v, _ = dec.raw_decode(obj[m.start():])
            except ValueError:
                continue
            f = find(v, depth + 1, want)
            if f: return f
    return None


def load(path):
    d = find(open(path, encoding='utf-8').read())
    if not d: sys.exit('could not find the report object (with "today" and "rows") in ' + path)
    return d


def n(d, k):
    return int((d or {}).get(k) or 0)


def yes(d, k):
    return 'Yes' if (d or {}).get(k) in (True, 'true') else 'No'


def hdr(ws, row, last):
    for c in range(1, last + 1):
        x = ws.cell(row, c); x.font = FH; x.fill = HF; x.alignment = C; x.border = B


TOOL_NAMES = {'course_search': 'Course-page search', 'page_reads': 'Course page reads', 'ai_page_match': 'AI page match', 'ai_field_reader': 'AI field reader',
              'fee_reader': 'Fee reader (AI)', 'scholarship_ai': 'Scholarship AI', 'other_credits': 'Other scraper credits', 'flags_raised': 'Flags raised'}


def toolset_sheet(wb, d):
    """22 Oct 2026 (Platform Admin "keep the stat how much of toolset we used used successfully and failed"): per provider and tool."""
    tools = d.get('tools') or []
    if not tools: return
    ts = wb.create_sheet('Toolset use')
    order = list(TOOL_NAMES)
    tools = sorted(tools, key=lambda t: (html.unescape(t.get('name') or '').lower(), order.index(t['tool']) if t['tool'] in order else 99))
    ts.cell(1, 1, f"Toolset use by {COUNTRY.get(CC, CC)} providers: {d.get('tools_day') or ''} and the 7 days to it").font = FB
    h = ['Tool', 'Used (day)', 'Worked (day)', 'Did not work (day)', 'Credits (day)', 'AI US$ (day)', 'Used (7 days)', 'Worked (7 days)', 'Did not work (7 days)', 'Credits (7 days)', 'AI US$ (7 days)', 'Worked share (7 days)']
    for j, t in enumerate(h, 1): ts.cell(3, j, t)
    hdr(ts, 3, len(h))
    first = 4 + len(order) + 3; last = first + len(tools) - 1
    for i, key in enumerate(order, 4):
        ts.cell(i, 1, TOOL_NAMES[key])
        for j in range(2, 12):
            ts.cell(i, j, f'=SUMIF($B${first}:$B${last},$A{i},{L(j + 1)}${first}:{L(j + 1)}${last})')
        ts.cell(i, 12, f'=IF(G{i}=0,"",H{i}/G{i})'); ts.cell(i, 12).number_format = '0%'
        for j in range(1, 13): ts.cell(i, j).font = F; ts.cell(i, j).border = B
        for j in (6, 11): ts.cell(i, j).number_format = '0.000'
    h2 = ['Provider', 'Tool'] + h[1:]
    for j, t in enumerate(h2, 1): ts.cell(first - 1, j, t)
    hdr(ts, first - 1, len(h2))
    for i, t in enumerate(tools, first):
        vals = [html.unescape(t.get('name') or ''), TOOL_NAMES.get(t['tool'], t['tool']), t.get('used') or 0, t.get('ok') or 0, t.get('failed') or 0, float(t.get('credits') or 0), float(t.get('cost') or 0),
                t.get('w_used') or 0, t.get('w_ok') or 0, t.get('w_failed') or 0, float(t.get('w_credits') or 0), float(t.get('w_cost') or 0)]
        for j, v in enumerate(vals, 1): ts.cell(i, j, v).font = F
        ts.cell(i, 13, f'=IF(H{i}=0,"",I{i}/H{i})').number_format = '0%'
        for j in (7, 12): ts.cell(i, j).number_format = '0.000'
    ts.cell(first - 1, 13, 'Worked share (7 days)'); hdr(ts, first - 1, 13)
    ts.column_dimensions['A'].width = 44; ts.column_dimensions['B'].width = 22
    for c in range(3, 14): ts.column_dimensions[L(c)].width = 12
    ts.freeze_panes = ts.cell(first, 1)
    ts.auto_filter.ref = f'A{first - 1}:M{last}'



def resolutions_sheet(wb, d):
    """22 Oct 2026 (Platform Admin 13:28: "Include resolution to be programmed after everyday report with lesson learnt")."""
    path = os.path.join(BASE_DIR, f'{CC.lower()}-resolutions.json')
    if not os.path.exists(path): return
    r = json.load(open(path, encoding='utf-8')); m = d.get('measures') or {}
    ws = wb.create_sheet('Resolutions & lessons', 1)
    ws.cell(1, 1, f"Resolutions programmed after the {CC} enrichment review, with yesterday's figures ({m.get('day', 'no figures yet')})").font = FB
    h = ['#', 'Gap', 'Resolution', 'Status', 'Release', 'Measure', 'Yesterday / now']
    for j, t in enumerate(h, 1): ws.cell(3, j, t)
    hdr(ws, 3, len(h))
    i = 4
    for x in r['resolutions']:
        ms = x.get('measures') or [['', '']]
        first = i
        for k, (key, label) in enumerate(ms):
            if k == 0:
                for j, v in enumerate([x['n'], x['gap'], x['resolution'], x['status'], x.get('release', '')], 1): ws.cell(i, j, v)
            ws.cell(i, 6, label); ws.cell(i, 7, m.get(key, '—') if key else '—')
            for j in range(1, 8):
                c = ws.cell(i, j); c.font = F; c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
            i += 1
        if i - 1 > first:
            for j in range(1, 6): ws.merge_cells(start_row=first, start_column=j, end_row=i - 1, end_column=j)
    i += 1
    ws.cell(i, 1, 'Lessons learnt').font = FB; i += 1
    for j, t in enumerate(['Date', 'Lesson', 'What changed'], 1): ws.cell(i, j, t)
    hdr(ws, i, 3); i += 1
    for x in r.get('lessons', []):
        for j, v in enumerate([x['date'], x['lesson'], x['action']], 1):
            c = ws.cell(i, j, v); c.font = F; c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
        i += 1
    for col, w in zip('ABCDEFG', (6, 40, 60, 10, 18, 42, 14)): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A4'


STATE = {'no_reader': 'no reader yet', 'not_on_page': 'not on the course page', 'site_known': 'course page not found yet', 'page_found': 'page found, not read yet',
         'in_review': 'in review', 'in_queue': 'in the queue', 'blocked': 'site blocks reading', 'no_website': 'no website', 'awaiting_l3': 'waiting for Layer 3',
         'candidate': 'value found, not admitted yet', 'missing_l1': 'not in the register'}
# Courses sheet: (heading, attribute whose state explains an empty cell, width); the order matches the course detail query's array
COURSE_COLS = [('Provider', None, 36), ('Course', None, 48), ('Course code', None, 14), ('Official course page', 'official_url', 40), ('Study level', 'level', 18),
               ('Field of study', 'field_of_study', 24), ('Duration', 'duration', 11), ('Delivery mode', 'delivery', 14), ('Campuses', 'campus', 30),
               ('Campus city', 'campus_city', 18), ('Open to international students', 'international', 14), ('Open to domestic students', None, 12),
               ('Tuition fee', 'provider_tuition', 34), ('Registered tuition (CRICOS)', 'registered_tuition', 22), ('Intakes / start dates', 'intakes', 30),
               ('Application deadline', 'deadline', 16), ('English requirements', 'english', 28), ('Academic entry requirements', 'academic_entry', 30),
               ('Entry requirements (course)', None, 40), ('Course description', 'description', 50)]
ENTRY_COLS = [('Provider', 36), ('Kind', 16), ('Status', 11), ('Level', 20), ('Study level', 16), ('Applicant country', 10), ('Qualification', 26), ('Benchmark band', 14),
              ('Minimum grade', 14), ('Minimum value', 10), ('Grade scale', 14), ('Audience', 14), ('Intake year', 9), ('Edition', 12), ('Quote', 50), ('Source page', 40), ('As at', 11)]
SCHOL_COLS = [('Provider', 36), ('Scholarship', 40), ('Type', 18), ('Audience', 18), ('Award (as written)', 30), ('Award type', 14), ('Amount', 11), ('Currency', 9),
              ('Percentage', 10), ('Up to (maximum)', 10), ('Applies to fee', 14), ('Fee basis', 12), ('Duration basis', 12), ('Application required', 11),
              ('Applications open', 12), ('Applications close', 12), ('Academic year', 9), ('Study start', 14), ('Nationalities', 20), ('Description', 50),
              ('Source page', 40), ('Publication status', 12)]


def cell_value(v):
    if v is None or v == '': return None
    if v is True or v == 'true': return 'Yes'
    if v is False or v == 'false': return 'No'
    if isinstance(v, str): return ILLEGAL_CHARACTERS_RE.sub('', html.unescape(v))[:32000]
    return v


def load_detail(paths, cc):
    """Join the parts of the course detail query; every part 0..parts-1 must be there, all for the same country."""
    parts = {}
    for p in paths:
        x = find(open(p, encoding='utf-8').read(), want=lambda o: o.get('kind') == 'course_detail')
        if not x: sys.exit('could not find the course detail object (kind "course_detail") in ' + p)
        if (x.get('country') or '').upper() != cc: sys.exit(f'{p} is for {x.get("country")}, not {cc}')
        parts[int(x.get('part') or 0)] = x
    want = int(next(iter(parts.values())).get('parts') or 1)
    if sorted(parts) != list(range(want)): sys.exit(f'course detail parts found {sorted(parts)}, expected 0 to {want - 1}')
    out = {'courses': [], 'entry': [], 'scholarships': [], 'day': None, 'coverage_at': None}
    for k in sorted(parts):
        for key in ('courses', 'entry', 'scholarships'): out[key] += parts[k].get(key) or []
        out['day'] = parts[k].get('day'); out['coverage_at'] = max(out['coverage_at'] or '', parts[k].get('coverage_at') or '')
    srt = lambda r: (html.unescape(r[0] or '').lower(), html.unescape(str(r[1] or '')).lower())
    for key in ('courses', 'entry', 'scholarships'): out[key].sort(key=srt)
    return out


def table(ws, cols, rows, conv):
    for j, (h, w) in enumerate(cols, 1):
        ws.cell(1, j, h); ws.column_dimensions[L(j)].width = w
    hdr(ws, 1, len(cols))
    for row in rows: ws.append(conv(row))
    ws.freeze_panes = 'A2'  # rows only (phone viewers)
    ws.auto_filter.ref = f'A1:{L(len(cols))}{max(len(rows) + 1, 2)}'
    ws.row_dimensions[1].height = 42


def detail_sheets(wb, det, n_schol):
    """31 Oct 2026: one row per active course with each course field value; an empty field shows why it is not held."""
    cs = wb.create_sheet('Courses')
    cols = [(h, w) for h, _, w in COURSE_COLS] + [('Provider scholarships (active)', 12)]

    def course(r):
        states = r[len(COURSE_COLS)] if len(r) > len(COURSE_COLS) else None
        out = []
        for j, (h, attr, _) in enumerate(COURSE_COLS):
            v = cell_value(r[j])
            if v is None and attr and states and attr in states: v = '— ' + STATE.get(states[attr], states[attr].replace('_', ' '))
            out.append(v)
        out.append(n_schol.get(html.unescape(r[0] or ''), 0))
        return out
    table(cs, cols, det['courses'], course)
    table(wb.create_sheet('Entry requirements'), ENTRY_COLS, det['entry'], lambda r: [cell_value(v) for v in r])
    table(wb.create_sheet('Scholarships'), SCHOL_COLS, det['scholarships'], lambda r: [cell_value(v) for v in r])


def main():
    global CC
    d = load(sys.argv[1])
    CC = (d.get('country') or 'NZ').upper(); name = COUNTRY.get(CC, CC)
    today, prev, base = d['today'], d.get('previous') or '', d.get('baseline') or '2026-10-20'
    ROOT = os.path.join(BASE_DIR, f'{CC.lower()}-fields')
    rows = sorted(d['rows'], key=lambda r: (-int(r.get('courses') or 0), html.unescape(r.get('name') or '').lower()))
    det = load_detail(sys.argv[2:], CC) if len(sys.argv) > 2 else None
    os.makedirs(ROOT, exist_ok=True)
    out = os.path.join(ROOT, f'{CC}-course-data-{today}.xlsx')
    wb = Workbook()

    r = wb.active; r.title = 'Read me'
    start = (f'Start: values held at the start of 21 Oct 2026, before the NZ enrichment run ({base}; worked out from what was filled that day).' if CC == 'NZ'
             else f'Start: values held on {base}, the first day recorded for {name}.')
    notes = [(f'{name} course data by provider: {today}', FB), ('', F),
             (start, F),
             (f'Previous: the day before ({prev or "none yet"}). Today: {today}. Each day is recorded at about 6:55 am Melbourne time.', F),
             (f'Scope: {len(rows)} {CC} providers with active courses, {sum(int(x.get("courses") or 0) for x in rows):,} active courses.', F), ('', F),
             ('Sheets', FB),
             (f'{CC} summary: each field across {name} at the start, the previous day and today, with the changes and % filled.', F),
             ('Per provider: one row per provider; for each course field, the number of courses holding it at the start, the previous day and today, the change today and the change since the start. Provider fields show Yes or No.', F),
             ('Long list: the same figures as one row per provider and field, for filtering and pivot tables.', F),
             *([] if CC != 'NZ' else [('Resolutions & lessons: each gap found in the NZ enrichment review, what was programmed for it, its status and one or more figures from yesterday that show whether it works; then the lessons learnt.', F)]),
             *([] if not det else [
                 (f'Courses: one row per active course ({len(det["courses"]):,}) with the value held for each course field. An empty field shows why it is not held '
                  '(for example "— not on the course page" or "— no reader yet"). ' f'Values as at {det["day"]}; course states from {(det["coverage_at"] or "")[:16].replace("T", " ")} UTC.', F),
                 ('Entry requirements: the entry rules each provider publishes (minimum results, band ladders), active and proposed; course-level entry requirements are in the Courses sheet.', F),
                 (f'Scholarships: each active scholarship of these providers ({len(det["scholarships"]):,}) with its award, dates, audience and source page. Scholarships are held per provider, not per course; the Courses sheet shows how many each provider has.', F)]),
             ('Toolset use: each tool per provider for the last recorded day and the last 7 days: times used, worked, did not work, scraper credits and AI US$, with a summary by tool at the top. "Worked" means: a correct course page (search, page match), a page read, a field filled, a fee page with fee rows, a scholarship applied.', F), ('', F),
             ('Notes', FB),
             ('Application deadline, academic entry requirements and course description have no reader yet.', F),
             *([('Six provider records are named "New Zealand Institute of Skills and Technology"; the website column tells them apart.', F)] if CC == 'NZ' else []),
             ('A provider with no figure for a day had no active courses recorded that day.', F)]
    for i, (t, f) in enumerate(notes, 1):
        c = r.cell(i, 1, t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
    r.column_dimensions['A'].width = 120

    ws = wb.create_sheet('Per provider')
    for j, h in enumerate(['Provider', 'Website', 'Courses'], 1):
        ws.cell(1, j, h); ws.merge_cells(start_row=1, start_column=j, end_row=2, end_column=j)
    col, cfcol = 4, {}
    for key, lab in CF:
        ws.cell(1, col, lab); ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 4)
        for k, s in enumerate(['Start', 'Previous', 'Today', 'Change today', 'Change since start']): ws.cell(2, col + k, s)
        cfcol[key] = col; col += 5
    pcol = col
    for i, (key, lab) in enumerate(PF):
        c = pcol + 3 * i; ws.cell(1, c, lab); ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + 2)
        for k, s in enumerate(['Start', 'Previous', 'Today']): ws.cell(2, c + k, s)
    last = pcol + 3 * len(PF) - 1
    hdr(ws, 1, last); hdr(ws, 2, last)
    for i, x in enumerate(rows, 3):
        ws.cell(i, 1, html.unescape(x.get('name') or '')); ws.cell(i, 2, x.get('website') or '—'); ws.cell(i, 3, int(x.get('courses') or 0))
        for key, _ in CF:
            c = cfcol[key]
            ws.cell(i, c, n(x.get('b'), key)); ws.cell(i, c + 1, n(x.get('y'), key)); ws.cell(i, c + 2, n(x.get('t'), key))
            ws.cell(i, c + 3, f'={L(c + 2)}{i}-{L(c + 1)}{i}'); ws.cell(i, c + 4, f'={L(c + 2)}{i}-{L(c)}{i}')
            ws.cell(i, c + 2).fill = SF
            for k in (3, 4): ws.cell(i, c + k).number_format = '+0;-0;-'
        for k, (key, _) in enumerate(PF):
            c = pcol + 3 * k
            ws.cell(i, c, yes(x.get('bp'), key)); ws.cell(i, c + 1, yes(x.get('yp'), key)); ws.cell(i, c + 2, yes(x.get('tp'), key))
        for c in range(1, last + 1):
            ws.cell(i, c).font = F; ws.cell(i, c).border = B
            if c >= 3: ws.cell(i, c).alignment = MID
    nrow = len(rows) + 2
    ws.column_dimensions['A'].width = 48; ws.column_dimensions['B'].width = 30; ws.column_dimensions['C'].width = 9
    for c in range(4, last + 1): ws.column_dimensions[L(c)].width = 10
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = 'A3'  # rows only: frozen columns hide the provider name in some phone viewers
    ws.auto_filter.ref = f'A2:{L(last)}{nrow}'

    sm = wb.create_sheet(f'{CC} summary', 1)
    h = ['Field', 'Level', 'Total', 'Start', 'Previous', 'Today', 'Change today', 'Change since start', '% start', '% today']
    for j, t in enumerate(h, 1): sm.cell(1, j, t)
    hdr(sm, 1, len(h))
    i = 2
    rng = lambda c: f"'Per provider'!{L(c)}3:{L(c)}{nrow}"
    for key, lab in CF:
        c = cfcol[key]
        sm.append([lab, 'Course', f"=SUM({rng(3)})", f'=SUM({rng(c)})', f'=SUM({rng(c + 1)})', f'=SUM({rng(c + 2)})',
                   f'=F{i}-E{i}', f'=F{i}-D{i}', f'=IF(C{i}=0,0,D{i}/C{i})', f'=IF(C{i}=0,0,F{i}/C{i})']); i += 1
    for k, (key, lab) in enumerate(PF):
        c = pcol + 3 * k
        sm.append([lab, 'Provider', f"=COUNTA({rng(1)})", f'=COUNTIF({rng(c)},"Yes")', f'=COUNTIF({rng(c + 1)},"Yes")', f'=COUNTIF({rng(c + 2)},"Yes")',
                   f'=F{i}-E{i}', f'=F{i}-D{i}', f'=IF(C{i}=0,0,D{i}/C{i})', f'=IF(C{i}=0,0,F{i}/C{i})']); i += 1
    for rr in range(2, i):
        for c in range(1, len(h) + 1):
            x = sm.cell(rr, c); x.font = F; x.border = B
            if c >= 3: x.alignment = MID
        for c in (3, 4, 5, 6): sm.cell(rr, c).number_format = '#,##0'
        for c in (7, 8): sm.cell(rr, c).number_format = '+#,##0;-#,##0;-'
        for c in (9, 10): sm.cell(rr, c).number_format = '0.0%'
        sm.cell(rr, 6).fill = SF
    sm.cell(i + 1, 1, 'Course fields count courses; provider fields count providers. Figures come from the Per provider sheet.').font = Font(name='Arial', size=9, italic=True)
    sm.column_dimensions['A'].width = 32
    for c in 'BCDEFGHIJ': sm.column_dimensions[c].width = 13
    sm.freeze_panes = 'A2'

    lg = wb.create_sheet('Long list')
    lg.append(['Provider', 'Website', 'Courses', 'Level', 'Field', 'Start', 'Previous', 'Today', 'Change today', 'Change since start'])
    hdr(lg, 1, 10)
    i = 2
    for x in rows:
        name = html.unescape(x.get('name') or ''); web = x.get('website') or '—'; cs = int(x.get('courses') or 0)
        for key, lab in CF:
            lg.append([name, web, cs, 'Course', lab, n(x.get('b'), key), n(x.get('y'), key), n(x.get('t'), key), f'=H{i}-G{i}', f'=H{i}-F{i}']); i += 1
        for key, lab in PF:
            v = lambda dd: 1 if (dd or {}).get(key) in (True, 'true') else 0
            lg.append([name, web, cs, 'Provider (1 = held)', lab, v(x.get('bp')), v(x.get('yp')), v(x.get('tp')), f'=H{i}-G{i}', f'=H{i}-F{i}']); i += 1
    for row in lg.iter_rows(min_row=2):
        for x in row: x.font = F
    for c, w in (('A', 48), ('B', 30), ('D', 18), ('E', 30)): lg.column_dimensions[c].width = w
    lg.freeze_panes = 'A2'; lg.auto_filter.ref = f'A1:J{i - 1}'
    if det:
        ns = {}
        for x in det['scholarships']: k = html.unescape(x[0] or ''); ns[k] = ns.get(k, 0) + 1
        detail_sheets(wb, det, ns)
    toolset_sheet(wb, d)
    resolutions_sheet(wb, d)
    wb.save(out)

    # three-line summary for the run's report
    tot = lambda dk, key: sum(n(x.get(dk), key) for x in rows)
    moved = [(lab, tot('t', k) - tot('y', k), tot('t', k) - tot('b', k)) for k, lab in CF]
    print(out)
    print('Change today: ' + (', '.join(f'{lab} {a:+,}' for lab, a, _ in moved if a) or 'none'))
    print(f'Change since {base}: ' + (', '.join(f'{lab} {b:+,}' for lab, _, b in moved if b) or 'none'))
    if det: print(f'Course detail: {len(det["courses"]):,} courses, {len(det["entry"]):,} provider entry rules, {len(det["scholarships"]):,} scholarships')


if __name__ == '__main__':
    main()
